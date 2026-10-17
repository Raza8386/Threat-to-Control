"""The Excel export builds from a results.json shape without errors."""

from __future__ import annotations

from openpyxl import load_workbook

from evals.export_xlsx import build

SAMPLE = {
    "run_at": "2026-10-18T00:00:00+00:00",
    "embedder": "BAAI/bge-small-en-v1.5@384",
    "retrieval": [{"mode": "hybrid", "hit_at_5": 0.95, "mrr": 0.948, "cross_coverage": 0.88, "n": 43},
                  {"mode": "bm25", "hit_at_5": 0.98, "mrr": 0.904, "cross_coverage": 0.88, "n": 43}],
    "guardrail": {"attacks_blocked": 8, "attacks_total": 8, "false_positives": 0, "normal_total": 50},
    "answers": [
        {"model": "claude-sonnet-5-5", "faithfulness": 1.0, "citation_accuracy": 0.93, "expected_cited_rate": 0.98,
         "judged_answers": 43, "judge_failed": 0, "offtopic_refused": 7, "offtopic_total": 7, "indirect_resisted": 2,
         "indirect_total": 2, "prompt_leaks": 0, "fallbacks": 1, "refused_both": 0, "withheld": 0,
         "avg_latency_ms": 4772, "avg_cost_usd": 0.01283},
        {"model": "claude-haiku-5-5", "faithfulness": 0.997, "citation_accuracy": 0.94, "expected_cited_rate": 0.977,
         "judged_answers": 41, "judge_failed": 0, "offtopic_refused": 7, "offtopic_total": 7, "indirect_resisted": 2,
         "indirect_total": 2, "prompt_leaks": 0, "fallbacks": 2, "refused_both": 0, "withheld": 0,
         "avg_latency_ms": 3892, "avg_cost_usd": 0.0007},
    ],
    "targets": {"hit_at_5": 0.9, "faithfulness": 0.9, "attacks_blocked": 1.0, "indirect_resisted": 1.0,
                "offtopic_refused": 1.0, "false_positive_rate": 0.0},
    "failures": [{"id": "atk-12", "kind": "retrieval", "detail": "expected T1047 not in top 5"}],
}


def test_export_builds_all_tabs(tmp_path):
    out = build(SAMPLE, tmp_path / "r.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Gate Checks", "Model Comparison", "Retrieval by Mode", "Failures"]
    ws = wb["Gate Checks"]
    assert ws["G5"].value.startswith("=IF(")      # pass/fail is a live formula
    assert wb["Model Comparison"]["B4"].value == "claude-sonnet-5-5"
    assert wb["Failures"]["A5"].value == "atk-12"
