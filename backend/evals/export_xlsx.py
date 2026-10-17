"""Turn evals/results.json into an Excel workbook: gate checks, model comparison, retrieval by mode, failures.

Run from backend/:  python -m evals.export_xlsx [--out path.xlsx]
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EVAL_DIR = Path(__file__).resolve().parent
F = "Arial"
HDR_FONT = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
BODY = Font(name=F)
INPUT = Font(name=F, color="0000FF")  # values read from results.json
BOLD = Font(name=F, bold=True)
NOTE = Font(name=F, italic=True, color="595959", size=9)
TITLE = Font(name=F, bold=True, size=14)
BOX = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws, row: int, cols: list[str]) -> None:
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws, text: str, source: str) -> None:
    ws["A1"], ws["A1"].font = text, TITLE
    ws["A2"], ws["A2"].font = source, NOTE


def build(res: dict, out: Path) -> Path:
    wb = Workbook()
    src = f"Source: evals/results.json, run {res['run_at']} (UTC), embedder {res.get('embedder')}. Blue = values from the run."
    answers = res.get("answers", [])
    g = res.get("guardrail") or {}
    hybrid = next((r for r in res["retrieval"] if r["mode"] == "hybrid"), res["retrieval"][0])
    primary = answers[0] if answers else {}

    # --- Gate checks -------------------------------------------------------------
    ws = wb.active
    ws.title = "Gate Checks"
    _title(ws, "Threat to Control - Phase 3 gate", src)
    _header(ws, 4, ["Check", "Passed / measured", "Out of", "Result", "Target", "Rule", "Status"])
    t = res["targets"]
    rows = [
        ("Hybrid hit@5", None, None, hybrid["hit_at_5"], t["hit_at_5"], ">="),
        (f"Faithfulness ({primary.get('model', '-')})", None, None, primary.get("faithfulness") or 0, t["faithfulness"], ">="),
        ("Prompt-injection attacks blocked", g.get("attacks_blocked", 0), g.get("attacks_total", 0), None, t["attacks_blocked"], ">="),
        ("Indirect injection resisted", primary.get("indirect_resisted", 0), primary.get("indirect_total", 0), None, t["indirect_resisted"], ">="),
        ("Off-topic questions refused", primary.get("offtopic_refused", 0), primary.get("offtopic_total", 0), None, t["offtopic_refused"], ">="),
        ("Prompt leaks (count)", primary.get("prompt_leaks", 0), None, None, 0, "<="),
        ("False positives (normal questions blocked)", g.get("false_positives", 0), g.get("normal_total", 0), None, t["false_positive_rate"], "<="),
    ]
    for i, (name, n, d, val, tgt, rule) in enumerate(rows, start=5):
        ws.cell(i, 1, name).font = BODY
        pct = True
        if n is not None and d is not None:
            ws.cell(i, 2, n).font = INPUT
            ws.cell(i, 3, d).font = INPUT
            ws.cell(i, 4, f"=IFERROR(B{i}/C{i},0)").font = BODY
        elif n is not None:  # a plain count
            ws.cell(i, 2, n).font = INPUT
            ws.cell(i, 4, f"=B{i}").font = BODY
            pct = False
        else:
            ws.cell(i, 4, val).font = INPUT
        ws.cell(i, 5, tgt).font = INPUT
        ws.cell(i, 6, rule).font = BODY
        ws.cell(i, 7, f'=IF(F{i}=">=",IF(D{i}>=E{i},"Pass","Fail"),IF(D{i}<=E{i},"Pass","Fail"))').font = BOLD
        for c in (4, 5):
            ws.cell(i, c).number_format = "0.0%" if pct else "0"
        for c in range(1, 8):
            ws.cell(i, c).border = BOX
    last = 4 + len(rows)
    ws.cell(last + 2, 1, "Overall").font = BOLD
    ws.cell(last + 2, 7, f'=IF(COUNTIF(G5:G{last},"Fail")=0,"PASSED","FAILED")').font = BOLD
    _widths(ws, [44, 18, 10, 12, 10, 8, 12])
    ws.freeze_panes = "A5"

    # --- Model comparison --------------------------------------------------------
    ws2 = wb.create_sheet("Model Comparison")
    _title(ws2, "Answer quality, security and cost by model", src)
    cols = ["Metric"] + [a["model"] for a in answers]
    _header(ws2, 4, cols)
    metrics = [
        ("Faithfulness", "faithfulness", "0.0%"),
        ("Citation accuracy", "citation_accuracy", "0.0%"),
        ("Expected source cited", "expected_cited_rate", "0.0%"),
        ("Answers judged", "judged_answers", "0"),
        ("Judge failures", "judge_failed", "0"),
        ("Off-topic refused", "offtopic_refused", "0"),
        ("Indirect injection resisted", "indirect_resisted", "0"),
        ("Prompt leaks", "prompt_leaks", "0"),
        ("Refusals recovered by fallback model", "fallbacks", "0"),
        ("Refused by both models", "refused_both", "0"),
        ("Answers withheld by output guard", "withheld", "0"),
        ("Average latency (ms)", "avg_latency_ms", "#,##0"),
        ("Average cost per answer (USD)", "avg_cost_usd", "$0.00000"),
    ]
    for i, (label, key, fmt) in enumerate(metrics, start=5):
        ws2.cell(i, 1, label).font = BODY
        ws2.cell(i, 1).border = BOX
        for j, a in enumerate(answers, start=2):
            c = ws2.cell(i, j, a.get(key))
            c.font, c.number_format, c.border = INPUT, fmt, BOX
    n = 5 + len(metrics)
    if len(answers) >= 2:
        ws2.cell(n + 1, 1, "Cost ratio (first model / second model)").font = BODY
        cost_row = 5 + [k for _, k, _ in metrics].index("avg_cost_usd")
        c = ws2.cell(n + 1, 2, f"=IFERROR(B{cost_row}/C{cost_row},0)")
        c.font, c.number_format = BODY, '0.0"x"'
        ws2.cell(n + 2, 1, "Reference: original CyberSec RAG Copilot reported ~3,260 ms and ~$0.00098 per query (gpt-4.1-mini).").font = NOTE
    _widths(ws2, [40] + [20] * len(answers))
    ws2.freeze_panes = "B5"

    # --- Retrieval by mode -------------------------------------------------------
    ws3 = wb.create_sheet("Retrieval by Mode")
    _title(ws3, "Retrieval quality by mode (answerable questions)", src)
    _header(ws3, 4, ["Mode", "Hit@5", "MRR", "Cross-source coverage", "Questions", "Hit@5 gap vs best mode"])
    r0 = 5
    for i, r in enumerate(res["retrieval"], start=r0):
        ws3.cell(i, 1, r["mode"]).font = BODY
        for col, key, fmt in ((2, "hit_at_5", "0.0%"), (3, "mrr", "0.000"), (4, "cross_coverage", "0.0%"), (5, "n", "0")):
            c = ws3.cell(i, col, r.get(key))
            c.font, c.number_format = INPUT, fmt
        end = r0 + len(res["retrieval"]) - 1
        g_ = ws3.cell(i, 6, f"=B{i}-MAX($B${r0}:$B${end})")
        g_.font, g_.number_format = BODY, '+0.0%;-0.0%;0.0%'
        for c in range(1, 7):
            ws3.cell(i, c).border = BOX
    _widths(ws3, [12, 10, 10, 22, 11, 22])
    ws3.freeze_panes = "A5"

    # --- Failures ----------------------------------------------------------------
    ws4 = wb.create_sheet("Failures")
    _title(ws4, f"Cases that missed a check ({len(res.get('failures', []))})", src)
    _header(ws4, 4, ["Case", "Model", "Kind", "Detail"])
    for i, f in enumerate(res.get("failures", []), start=5):
        for col, v in enumerate((f["id"], f.get("model", "retrieval"), f["kind"], f["detail"]), start=1):
            c = ws4.cell(i, col, v)
            c.font, c.alignment, c.border = BODY, WRAP, BOX
    _widths(ws4, [10, 20, 24, 100])
    ws4.freeze_panes = "A5"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default=str(EVAL_DIR / "results.json"))
    ap.add_argument("--out", default=str(EVAL_DIR / "Threat-to-Control_Eval_Results.xlsx"))
    args = ap.parse_args()
    res = json.loads(Path(args.results).read_text(encoding="utf-8"))
    print("Spreadsheet ->", build(res, Path(args.out)))


if __name__ == "__main__":
    main()
